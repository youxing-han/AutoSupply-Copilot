"""Fictional multi-SKU planning input, XLSX interchange and graph execution."""
from __future__ import annotations

import io
import math
from zipfile import BadZipFile
from copy import deepcopy
from datetime import timedelta

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from workflow_runtime import run_planning_workflow
from batch_analysis import validate_batch_dataframe, REQUIRED_COLUMNS, SUPPLIER_STATUS_ALIASES
from inventory_planning import day, quantity, simulate_inventory
from supply_data import NORMAL_SUPPLY, generate_automotive_supply_data
from upload_safety import (validate_upload, validate_frame_size,
                           validate_worksheet_size, PLANNING_ROW_LIMITS)

SHEETS = {
    "settings": ("start_date", "horizon_days"),
    "parts": ("part_code", "part_name", "supplier", "current_stock"),
    "suppliers": ("supplier", "status", "qualified"),
    "demands": ("part_code", "date", "quantity"),
    "orders": ("po_id", "part_code", "supplier", "open_quantity", "arrival_date", "status"),
}
DISPLAY_SHEETS = {"parts": "SKU_Master", "demands": "Demand_Plan", "orders": "Purchase_Orders",
                  "suppliers": "Suppliers", "settings": "Settings"}
INTERVIEW_SCENARIOS = ("正常到货", "供应商交付延迟", "生产需求激增", "采购订单待确认")
ORDER_EXTRA_COLUMNS = ("original_arrival_date", "delivery_status")
PLANNING_DATA_VERSION = 5
PART_OPTIONAL = ("safety_days", "criticality", "backup_supplier", "backup_supplier_available", "destination_line", "unit_cost")


def optional_lead(value):
    if value is None or value == '' or (isinstance(value, float) and math.isnan(value)):
        return None
    return quantity(value, '标准交期')


class PlanningInputError(ValueError):
    """Invalid planning data, distinct from internal graph errors."""


def build_interview_demo(start_date="2026-10-12", horizon_days=30, scenario="正常到货") -> dict:
    """Four controlled scenarios from one fictional five-part baseline.

    SEN has only the 500-unit PO, independent of horizon; extending to 30 days
    therefore exposes later shortages. Other parts share unchanged PO schedules.
    """
    if scenario not in INTERVIEW_SCENARIOS:
        raise PlanningInputError("请选择四种内置演示情景之一。")
    start, horizon = day(start_date), quantity(horizon_days, "horizon_days")
    if not 1 <= horizon <= 90:
        raise PlanningInputError("预测周期必须为 1–90 天。")
    parts = generate_automotive_supply_data(NORMAL_SUPPLY).to_dict("records")
    parts[0]["criticality"] = "HIGH"
    demands = [{"part_code": p["part_code"], "date": (start + timedelta(days=i)).isoformat(),
                "quantity": int(p["daily_usage"])} for p in parts for i in range(horizon)]
    schedule = [(7, 500), (11, 1800), (8, 900), (14, 600), (14, 1000)]
    orders = []
    for p, (offset, q) in zip(parts, schedule):
        arrival = (start + timedelta(days=offset)).isoformat()
        orders.append({"po_id": f"SIM-PO-{p['part_code'].split('-')[0]}-01", "part_code": p["part_code"],
                       "supplier": p["supplier"], "open_quantity": q, "arrival_date": arrival,
                       "original_arrival_date": arrival, "status": "CONFIRMED", "delivery_status": "ON_TIME"})
    # An unchanged pending secondary order illustrates commitments versus options.
    arrival = (start + timedelta(days=24)).isoformat()
    orders.append({"po_id": "SIM-PO-BRG-02", "part_code": "BRG-002", "supplier": parts[1]["supplier"],
                   "open_quantity": 300, "arrival_date": arrival, "original_arrival_date": arrival,
                   "status": "PENDING", "delivery_status": "UNCONFIRMED"})
    if scenario == "供应商交付延迟":
        orders[0]["arrival_date"] = (start + timedelta(days=11)).isoformat()
        orders[0]["delivery_status"] = "DELAYED"
        parts[0]["supplier_status"] = "DELAYED"
    elif scenario == "生产需求激增":
        for row in demands:
            if row["part_code"] == "SEN-001" and 6 <= (day(row["date"]) - start).days <= 9:
                row["quantity"] = 160
    elif scenario == "采购订单待确认":
        orders[0].update(status="PENDING", delivery_status="UNCONFIRMED")
    suppliers = [{"supplier": p["supplier"], "lead_time_days": p["lead_time_days"],
                  "status": p["supplier_status"], "qualified": True} for p in parts]
    return {"settings": [{"start_date": start.isoformat(), "horizon_days": horizon}],
            "parts": parts, "suppliers": suppliers, "demands": demands, "orders": orders}


def normal_arrival_baseline(bundle: dict) -> dict:
    """Counterfactual keeps current stock/demand/quantities; restores PO dates.

    Pending is assumed confirmed only in this explicitly labeled comparison.
    It never replaces the actual forecast or asserts receipt feasibility.
    """
    baseline = validate_planning_bundle(bundle)
    for po in baseline["orders"]:
        if po["status"] != "CANCELLED":
            po["arrival_date"] = po["original_arrival_date"]
            po["status"] = "CONFIRMED"
            po["delivery_status"] = "ON_TIME"
    for p in baseline["parts"]:
        p["supplier_status"] = "NORMAL"
    for s in baseline["suppliers"]:
        s["status"] = "NORMAL"
    return validate_planning_bundle(baseline)


def build_planning_demo(start_date="2026-10-12", horizon_days=30, scenario="延迟到货") -> dict:
    start = day(start_date)
    horizon = quantity(horizon_days, "horizon_days")
    if not 1 <= horizon <= 90:
        raise ValueError("horizon_days must be between 1 and 90")
    parts = generate_automotive_supply_data(NORMAL_SUPPLY).iloc[[0, 3]].to_dict("records")
    parts[0]["criticality"] = "HIGH"
    parts[0]["supplier_status"] = "DELAYED" if scenario == "延迟到货" else "NORMAL"
    parts[1]["current_stock"] = 100
    parts[1]["daily_usage"] = 20
    parts[1]["criticality"] = "HIGH"
    suppliers = [{"supplier": p["supplier"], "lead_time_days": p["lead_time_days"],
                  "status": p["supplier_status"], "qualified": True} for p in parts]
    demands = [{"part_code": p["part_code"], "date": (start + timedelta(days=i)).isoformat(),
                "quantity": (40 if i % 7 == 4 else 0 if i % 7 == 6 else 20)
                if p["part_code"] == "BRK-004" else (120 if scenario == "波动需求" and i % 3 == 0 else 80)}
               for p in parts for i in range(horizon)]
    arrival = 11 if scenario == "延迟到货" else 7
    orders = [{"po_id": "SIM-PO-SEN-01", "part_code": "SEN-001", "supplier": parts[0]["supplier"],
               "open_quantity": 500, "arrival_date": (start + timedelta(days=arrival)).isoformat(),
               "status": "PENDING" if scenario == "待确认订单" else "CONFIRMED"},
              {"po_id": "SIM-PO-SEN-02", "part_code": "SEN-001", "supplier": parts[0]["supplier"],
               "open_quantity": 1500, "arrival_date": (start + timedelta(days=15)).isoformat(), "status": "CONFIRMED"},
              {"po_id": "SIM-PO-BRK-01", "part_code": "BRK-004", "supplier": parts[1]["supplier"],
               "open_quantity": 600, "arrival_date": (start + timedelta(days=7)).isoformat(), "status": "CONFIRMED"}]
    if scenario == "无采购订单":
        orders = []
    return {"settings": [{"start_date": start.isoformat(), "horizon_days": horizon}],
            "parts": parts, "suppliers": suppliers, "demands": demands, "orders": orders}


def _bool(value):
    if isinstance(value, bool):
        return value
    if str(value).lower() in ("true","是"):
        return True
    if str(value).lower() in ("false","否"):
        return False
    raise ValueError("qualified must be True or False")


def validate_planning_bundle(bundle: dict) -> dict:
    try:
        return _validate_planning_bundle(bundle)
    except (ValueError, KeyError, TypeError, OverflowError) as exc:
        raise PlanningInputError(str(exc)) from exc


def _validate_planning_bundle(bundle: dict) -> dict:
    if not isinstance(bundle, dict) or any(k not in bundle for k in SHEETS):
        raise ValueError("required sheets: " + ", ".join(SHEETS))
    for name, limit in PLANNING_ROW_LIMITS.items():
        if not isinstance(bundle[name], list) or len(bundle[name]) > limit:
            raise ValueError(f"{name}: 最多 {limit} 行数据。")
        validate_frame_size(pd.DataFrame(bundle[name]), limit)
    normalized = deepcopy(bundle)
    for row in normalized.get('suppliers',[]):
        row['status']=SUPPLIER_STATUS_ALIASES.get(row.get('status'),row.get('status'))
    for row in normalized.get('orders',[]):
        if 'status' in row:row['status']={'已确认':'CONFIRMED','待确认':'PENDING','已取消':'CANCELLED'}.get(row['status'],row['status'])
        if 'delivery_status' in row:row['delivery_status']={'按期到货':'ON_TIME','交付延迟':'DELAYED','待确认':'UNCONFIRMED','已取消':'CANCELLED'}.get(row['delivery_status'],row['delivery_status'])
    for name, columns in SHEETS.items():
        rows = normalized[name]
        if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
            raise ValueError(f"{name}: 必须为数据行列表。")
        required = columns
        for i, row in enumerate(rows, 2):
            missing = set(required) - set(row)
            if missing:
                raise ValueError(f"{name} 第 {i} 行缺少字段：{', '.join(sorted(missing))}")
    settings = normalized["settings"]
    if len(settings) != 1:
        raise ValueError("settings must contain exactly one row")
    start = day(settings[0]["start_date"], "start_date")
    horizon = quantity(settings[0]["horizon_days"], "horizon_days")
    if not 1 <= horizon <= 90:
        raise ValueError("horizon_days must be between 1 and 90")
    if not 1 <= len(normalized['parts']) <= 50:
        raise ValueError('请提供1—50种物料')
    seen = set()
    defaults = dict(safety_days=3, criticality='MEDIUM', backup_supplier=None,
                    backup_supplier_available=False, destination_line='', unit_cost=0)
    for p in normalized['parts']:
        for field in ('part_code', 'part_name', 'supplier'):
            if not isinstance(p[field], str) or not p[field].strip():
                raise ValueError(f'{field} 不能为空')
            p[field] = p[field].strip()
        if p['part_code'] in seen:
            raise ValueError('物料编码重复')
        seen.add(p['part_code'])
        for k, v in defaults.items():
            if k not in p or p[k] is None or (isinstance(p[k], float) and math.isnan(p[k])):
                p[k] = v
        p['criticality'] = {'高':'HIGH','中':'MEDIUM','低':'LOW'}.get(p['criticality'], p['criticality'])
        if p['criticality'] not in {'HIGH','MEDIUM','LOW'}:
            raise ValueError('物料关键程度必须为高、中或低')
        for field in ('safety_days','unit_cost'):
            value = p[field]
            if isinstance(value, bool) or not isinstance(value, (int,float)) or not math.isfinite(value) or value < 0:
                raise ValueError(f'{field} 必须为有限非负数')
        p['current_stock'] = quantity(p['current_stock'], 'current_stock')
        p['backup_supplier_available'] = _bool(p['backup_supplier_available'])
        p['lead_time_days'] = optional_lead(p.get('lead_time_days'))
    suppliers = {}
    for row in normalized["suppliers"]:
        name = row["supplier"]
        if not isinstance(name, str) or not name.strip() or name != name.strip() or name in suppliers:
            raise ValueError("supplier name is empty or duplicate")
        row['lead_time_days'] = optional_lead(row.get('lead_time_days'))
        if row["status"] not in {"NORMAL", "DELAYED", "AT_RISK", "DISRUPTED", "STOPPED", "UNAVAILABLE"}:
            raise ValueError("invalid supplier status")
        row["qualified"] = _bool(row["qualified"])
        suppliers[name] = row
    sku_set = {p["part_code"] for p in normalized["parts"]}
    for p in normalized["parts"]:
        quantity(p["current_stock"], "current_stock")
        if p["supplier"] not in suppliers:
            raise ValueError("part references unknown supplier")
        supplier = suppliers[p["supplier"]]
        supplied_status = SUPPLIER_STATUS_ALIASES.get(p.get('supplier_status',supplier['status']),p.get('supplier_status',supplier['status']))
        if supplied_status != supplier['status'] or (p['lead_time_days'] is not None and supplier['lead_time_days'] is not None and p['lead_time_days'] != supplier['lead_time_days']):
            raise ValueError("part and supplier status/lead time must agree")
        p['supplier_status'] = supplier['status']
        p['lead_time_days'] = p['lead_time_days'] if p['lead_time_days'] is not None else supplier['lead_time_days']
        p["backup_supplier"] = p["backup_supplier"] or None
    for row in normalized["demands"] + normalized["orders"]:
        if row["part_code"] not in sku_set:
            raise ValueError("demand/PO references unknown SKU")
    po_ids = set()
    for po in normalized["orders"]:
        original = day(po.get("original_arrival_date", po["arrival_date"]), "original_arrival_date")
        actual = day(po["arrival_date"], "arrival_date")
        delivery = po.get("delivery_status")
        if delivery is not None and delivery not in {"ON_TIME", "DELAYED", "UNCONFIRMED", "CANCELLED"}:
            raise ValueError("delivery_status 必须为 ON_TIME/DELAYED/UNCONFIRMED/CANCELLED")
        po["original_arrival_date"] = original.isoformat()
        po["arrival_date"] = actual.isoformat()
        po["delivery_status"] = ("UNCONFIRMED" if po["status"] == "PENDING" else "CANCELLED" if po["status"] == "CANCELLED"
                                 else "DELAYED" if actual > original else "ON_TIME")
        if po["po_id"] in po_ids:
            raise ValueError("duplicate PO identifier; each PO may appear only once")
        po_ids.add(po["po_id"])
        if po["supplier"] not in suppliers:
            raise ValueError("PO references unknown supplier")
        if po["status"] == "CONFIRMED" and not suppliers[po["supplier"]]["qualified"]:
            raise ValueError("confirmed PO supplier is not qualified")
    for p in normalized["parts"]:
        simulate_inventory(p, [r for r in normalized["demands"] if r["part_code"] == p["part_code"]],
                           [r for r in normalized["orders"] if r["part_code"] == p["part_code"]], start, horizon)
    for row in normalized["demands"]:
        row["date"] = day(row["date"]).isoformat()
        row['quantity'] = quantity(row['quantity'], 'quantity')
    for p in normalized['parts']:
        p['daily_usage'] = sum(r['quantity'] for r in normalized['demands'] if r['part_code'] == p['part_code']) / horizon
    normalized["settings"] = [{"start_date": start.isoformat(), "horizon_days": horizon}]
    return normalized


def planning_xlsx_bytes(bundle: dict, *, named_sheets=False, blank=False) -> bytes:
    """A native import template; existing project openpyxl dependency only."""
    normalized = validate_planning_bundle(bundle)
    if blank:
        normalized = {name: rows if name == "settings" else [] for name, rows in normalized.items()}
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        names = list(DISPLAY_SHEETS) if named_sheets else list(SHEETS)
        for name in names:
            columns = SHEETS[name]
            if name == "parts":
                columns = (*SHEETS['parts'], *PART_OPTIONAL)
            elif name == 'suppliers':
                columns = (*columns, 'lead_time_days')
            elif name == "orders":
                columns = (*columns, *ORDER_EXTRA_COLUMNS)
            frame = pd.DataFrame(normalized[name])
            if columns:
                frame = frame.reindex(columns=columns)
            for column in ("start_date", "date", "arrival_date", "original_arrival_date"):
                if column in frame:
                    frame[column] = pd.to_datetime(frame[column])
            sheet_name = DISPLAY_SHEETS[name] if named_sheets else name
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="176B5B")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[1].height = 32
            for i, col in enumerate(frame.columns, 1):
                width = max([len(str(col)), *(len(str(v)) for v in frame[col])]) + 2
                sheet.column_dimensions[get_column_letter(i)].width = min(max(width, 14), 38)
                for cells in sheet.iter_rows(min_row=2, min_col=i, max_col=i):
                    if col in {"start_date", "date", "arrival_date", "original_arrival_date"}:
                        cells[0].number_format = "yyyy-mm-dd"
                    elif col == "unit_cost":
                        cells[0].number_format = "#,##0.00"
    return output.getvalue()


def read_planning_xlsx(content: bytes) -> dict:
    validate_upload(content, xlsx=True)
    try:
        with pd.ExcelFile(io.BytesIO(content), engine="openpyxl") as workbook:
            if len(workbook.sheet_names) > 5:
                raise ValueError("动态计划仅接受五个业务工作表。")
            mapped = {}
            for name in SHEETS:
                matches = [s for s in (name, DISPLAY_SHEETS[name]) if s in workbook.sheet_names]
                if len(matches) > 1:
                    raise PlanningInputError(f"{name}: 新旧命名工作表同时存在，无法确定数据来源。")
                if matches:
                    mapped[name] = matches[0]
            missing = set(SHEETS) - set(mapped)
            if missing:
                raise ValueError("missing sheets: " + ", ".join(sorted(missing)))
            bundle = {}
            for name, columns in SHEETS.items():
                validate_worksheet_size(workbook.book[mapped[name]], PLANNING_ROW_LIMITS[name])
                # Read raw headers before pandas can rename duplicates to .1.
                headers = [c.value for c in next(workbook.book[mapped[name]].iter_rows(min_row=1, max_row=1))]
                if len(headers) != len(set(headers)):
                    raise ValueError(f"{name}: duplicate column headers")
                frame = pd.read_excel(workbook, sheet_name=mapped[name])
                validate_frame_size(frame, PLANNING_ROW_LIMITS[name])
                if any(c not in frame for c in columns):
                    raise ValueError(f"{name}: missing columns")
                bundle[name] = frame.to_dict("records")
    except (BadZipFile, InvalidFileException, EOFError) as exc:
        raise ValueError("invalid or damaged XLSX workbook") from exc
    return validate_planning_bundle(bundle)


def analyze_planning_bundle(bundle: dict) -> dict:
    normalized = validate_planning_bundle(bundle)
    settings = normalized["settings"][0]
    details, rows = {}, []
    for p in normalized["parts"]:
        sku = p["part_code"]
        demands = [r for r in normalized["demands"] if r["part_code"] == sku]
        orders = [r for r in normalized["orders"] if r["part_code"] == sku]
        # The safety threshold uses average planned demand, not a legacy fixed rate.
        p["daily_usage"] = sum(r["quantity"] for r in demands) / settings["horizon_days"]
        p["safety_stock"] = p["daily_usage"] * p["safety_days"]
        result = run_planning_workflow(p, {**settings, "demands": demands, "orders": orders})
        details[sku] = result
        projection = result["inventory_projection"]
        rows.append({"SKU": sku, "名称": p["part_name"], "风险": result["supplier_risk"]["risk_level"],
                     "首次缺料": projection["first_shortage_date"] or "无",
                     "首次缺料量": projection["first_shortage_quantity"],
                     "累计未满足需求": projection["total_unmet_demand"],
                     "缺料天数": projection["shortage_days"], "期末可用库存": projection["final_available"]})
    summary = pd.DataFrame(rows)
    summary["_risk"] = summary["风险"].map({"HIGH": 0, "MEDIUM": 1, "LOW": 2})
    summary = summary.sort_values(["_risk", "累计未满足需求", "SKU"], ascending=[True, False, True]).drop(columns="_risk").reset_index(drop=True)
    return {"summary": summary, "details": details, "bundle": normalized}
