import copy
import io
import sys
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "agent"))
from agents import build_supply_chain_graph, create_initial_state
from assistant import build_assistant_context, answer_supply_chain_question, AssistantUnavailableError
from inventory_planning import simulate_inventory, evaluate_dynamic_risk, day
from planning_data import (build_planning_demo, validate_planning_bundle, analyze_planning_bundle,
                           planning_xlsx_bytes, read_planning_xlsx)
from planning_ui import inventory_figure
from report_pdf import build_pdf_report
from supply_data import get_scenario_part, NORMAL_SUPPLY

START = date(2026, 10, 12)


def fixture(horizon=14, arrival=7, status="CONFIRMED"):
    part = get_scenario_part(NORMAL_SUPPLY, "SEN-001")
    part["safety_stock"] = 240
    demands = [{"part_code": "SEN-001", "date": (START + timedelta(days=i)).isoformat(),
                "quantity": 80} for i in range(horizon)]
    orders = [] if arrival is None else [{"po_id": "SIM-01", "part_code": "SEN-001", "supplier": part["supplier"],
                "open_quantity": 500, "arrival_date": (START + timedelta(days=arrival)).isoformat(), "status": status}]
    return part, demands, orders


class InventoryPlanningTests(unittest.TestCase):
    def simulate(self, arrival=7, status="CONFIRMED", horizon=14):
        return simulate_inventory(*fixture(horizon, arrival, status), START, horizon)

    def test_scenario1_confirmed_receipt_day8(self):
        r = self.simulate()
        self.assertEqual((r["final_available"], r["total_unmet_demand"], r["first_shortage_date"]), (180, 0, None))
        self.assertEqual(r["daily"][7]["confirmed_receipts"], 500)
        self.assertEqual(r["unmet_avoided_by_confirmed_orders"], 320)

    def test_scenario2_no_order_shortage_day11(self):
        r = self.simulate(None)
        self.assertEqual(r["first_shortage_date"], "2026-10-22")
        self.assertEqual((r["first_shortage_quantity"], r["total_unmet_demand"], r["shortage_days"]), (80, 320, 4))
        self.assertEqual(r["risk_windows"], [{"start": "2026-10-22", "end": "2026-10-25", "days": 4, "quantity": 320}])
        self.assertTrue(all(d["closing_available"] >= 0 for d in r["daily"]))

    def test_scenario3_delay_until_day12_no_backorder(self):
        r = self.simulate(11)
        self.assertEqual((r["first_shortage_date"], r["total_unmet_demand"], r["shortage_days"]), ("2026-10-22", 80, 1))
        self.assertEqual(r["daily"][11]["closing_available"], 420)
        self.assertEqual(r["final_available"], 260)
        self.assertEqual(r["unmet_avoided_by_confirmed_orders"], 240)

    def test_scenario4_variable_demand_matches_manual_inventory(self):
        p, d, o = fixture(4, 1)
        for row, q in zip(d, [200, 0, 900, 300]):
            row["quantity"] = q
        r = simulate_inventory(p, d, o, START, 4)
        self.assertEqual([x["closing_available"] for x in r["daily"]], [600, 1100, 200, 0])
        self.assertEqual(r["total_unmet_demand"], 100)
        self.assertEqual(r["first_shortage_date"], "2026-10-15")

    def test_scenario5_pending_does_not_remove_shortage(self):
        r = self.simulate(7, "PENDING")
        self.assertEqual(r["total_unmet_demand"], 320)
        self.assertEqual(len(r["confirmed_orders"]), 0)
        self.assertEqual(len(r["excluded_orders"]), 1)

    def test_cancelled_order_is_excluded(self):
        self.assertEqual(self.simulate(7, "CANCELLED")["total_unmet_demand"], 320)

    def test_same_day_receipt_before_production(self):
        p, d, o = fixture(1, 0)
        p["current_stock"] = 0
        r = simulate_inventory(p, d, o, START, 1)
        self.assertEqual((r["final_available"], r["total_unmet_demand"]), (420, 0))

    def test_final_day_receipt_and_next_day_boundary(self):
        self.assertEqual(self.simulate(13)["total_unmet_demand"], 240)
        self.assertEqual(self.simulate(14)["total_unmet_demand"], 320)

    def test_zero_demand_does_not_create_shortage(self):
        p, d, o = fixture()
        for row in d:
            row["quantity"] = 0
        r = simulate_inventory(p, d, o, START, 14)
        self.assertEqual((r["final_available"], r["total_unmet_demand"]), (1300, 0))

    def test_calendar_boundary_leap_day(self):
        p, d, o = fixture(3, None)
        for i, row in enumerate(d):
            row["date"] = (date(2028, 2, 28) + timedelta(days=i)).isoformat()
        r = simulate_inventory(p, d, o, "2028-02-28", 3)
        self.assertEqual([x["date"] for x in r["daily"]], ["2028-02-28", "2028-02-29", "2028-03-01"])

    def test_shortage_windows_are_not_merged_across_satisfied_day(self):
        p, d, o = fixture(3, 1)
        p["current_stock"] = 0
        d[2]["quantity"] = 600
        r = simulate_inventory(p, d, o, START, 3)
        self.assertEqual([x["days"] for x in r["risk_windows"]], [1, 1])

    def test_scenario6_invalid_inputs_are_rejected(self):
        mutations = [lambda p,d,o: d.clear(), lambda p,d,o: d[0].update(date="2026-02-30"),
                     lambda p,d,o: o.append(copy.deepcopy(o[0])), lambda p,d,o: d[0].update(quantity=-1),
                     lambda p,d,o: o[0].update(open_quantity=-1), lambda p,d,o: d[0].update(part_code="WRONG"),
                     lambda p,d,o: o[0].update(part_code="WRONG"), lambda p,d,o: d.append(copy.deepcopy(d[0])),
                     lambda p,d,o: p.update(current_stock=-1), lambda p,d,o: o[0].update(status="UNKNOWN"),
                     lambda p,d,o: o[0].update(arrival_date="2026-10-11"),
                     lambda p,d,o: d[0].update(quantity=float("nan")),
                     lambda p,d,o: d[0].update(quantity=True), lambda p,d,o: d[0].update(quantity=1.5)]
        for mutation in mutations:
            p, d, o = fixture()
            mutation(p,d,o)
            with self.subTest(mutation=mutation), self.assertRaises((ValueError, TypeError)):
                simulate_inventory(p,d,o,START,14)

    def test_horizon_and_date_validation(self):
        for horizon in [0, 91, -1, True, 14.2]:
            with self.subTest(horizon=horizon), self.assertRaises(ValueError):
                simulate_inventory(*fixture(), START, horizon)
        for value in ["12/10/2026", "2026-10-12T00:00:00", None]:
            with self.subTest(date=value), self.assertRaises(ValueError):
                day(value)

    def test_risk_distinguishes_quantity_safety_and_supplier(self):
        p, d, o = fixture()
        r = simulate_inventory(p,d,o,START,14)
        risk = evaluate_dynamic_risk(p, r)
        self.assertFalse(risk["quantity_shortage"])
        self.assertTrue(risk["safety_stock_warning"])
        self.assertEqual(risk["risk_level"], "MEDIUM")
        p["current_stock"] = 5000
        p["supplier_status"] = "DISRUPTED"
        risk = evaluate_dynamic_risk(p, simulate_inventory(p,d,o,START,14))
        self.assertEqual(risk["risk_level"], "MEDIUM")
        self.assertFalse(risk["quantity_shortage"])
        self.assertTrue(risk["supplier_warning"])
        p["criticality"] = "HIGH"
        p["current_stock"] = 800
        risk = evaluate_dynamic_risk(p, simulate_inventory(p,d,[],START,14))
        self.assertEqual(risk["risk_level"], "HIGH")

    def test_every_day_conserves_inventory_and_demand(self):
        for arrival in [None, 0, 7, 11, 13, 14]:
            r = self.simulate(arrival)
            for row in r["daily"]:
                self.assertEqual(row["opening_available"] + row["confirmed_receipts"], row["satisfied_demand"] + row["closing_available"])
                self.assertEqual(row["planned_demand"], row["satisfied_demand"] + row["unmet_demand"])
            self.assertEqual(r["initial_available"] + sum(x["confirmed_receipts"] for x in r["daily"]),
                             r["final_available"] + sum(x["satisfied_demand"] for x in r["daily"]))


class PlanningIntegrationTests(unittest.TestCase):
    def test_corrupt_workbook_and_duplicate_raw_headers_rejected(self):
        from openpyxl import load_workbook
        for content in [b"", b"not a zip workbook", b"a" * 5_000_001]:
            with self.subTest(length=len(content)), self.assertRaises(ValueError):
                read_planning_xlsx(content)
        workbook = load_workbook(io.BytesIO(planning_xlsx_bytes(build_planning_demo())))
        workbook["demands"]["D1"] = "date"
        payload = io.BytesIO()
        workbook.save(payload)
        with self.assertRaisesRegex(ValueError, "duplicate"):
            read_planning_xlsx(payload.getvalue())
    def test_xlsx_import_round_trip_and_typed_dates(self):
        from openpyxl import load_workbook
        for scenario in ["延迟到货", "无采购订单"]:
            b = build_planning_demo(scenario=scenario)
            payload = planning_xlsx_bytes(b)
            loaded = read_planning_xlsx(payload)
            self.assertEqual(analyze_planning_bundle(b)["summary"].to_dict("records"), analyze_planning_bundle(loaded)["summary"].to_dict("records"))
            workbook = load_workbook(io.BytesIO(payload))
            self.assertEqual(set(workbook.sheetnames), {"settings", "parts", "suppliers", "demands", "orders"})
            self.assertTrue(workbook["demands"]["B2"].is_date)
            self.assertEqual(workbook["parts"].freeze_panes, "A2")

    def test_bundle_rejects_cross_reference_duplicates_and_missing_data(self):
        mutations = [lambda b:b["parts"].clear(), lambda b:b["suppliers"].clear(),
                     lambda b:b["orders"][0].update(supplier="UNKNOWN"),
                     lambda b:b["demands"][0].update(part_code="UNKNOWN"),
                     lambda b:b["orders"].append(copy.deepcopy(b["orders"][0])),
                     lambda b:b["suppliers"][0].update(qualified=False),
                     lambda b:b["suppliers"][0].update(status="STOPPED"),
                     lambda b:b["settings"].append(copy.deepcopy(b["settings"][0]))]
        for mutation in mutations:
            b = build_planning_demo()
            mutation(b)
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                validate_planning_bundle(b)

    def test_actual_four_node_graph_report_trace_and_pdf(self):
        analysis = analyze_planning_bundle(build_planning_demo())
        self.assertEqual(len(analysis["details"]), 2)
        self.assertEqual(analysis["summary"].iloc[0]["SKU"], "SEN-001")
        for r in analysis["details"].values():
            self.assertEqual([t["node_name"] for t in r["execution_trace"]], ["Inventory Monitor", "Supplier Risk", "Replenishment Decision", "Management Report"])
            self.assertTrue(all(t["execution_status"] == "SUCCESS" for t in r["execution_trace"]))
            self.assertIn("动态物料计划报告", r["management_report"])
            self.assertIn(str(r["inventory_projection"]["total_unmet_demand"]), r["management_report"])
            self.assertTrue(build_pdf_report(r).startswith(b"%PDF"))
            figure=inventory_figure(r)
            self.assertIn('当日确认到货量',[t.name for t in figure.data])
            self.assertEqual(list(next(t for t in figure.data if t.name=='当日计划需求').y),[d['planned_demand'] for d in r['inventory_projection']['daily']])

    def test_bad_planning_input_has_failed_trace_without_stale_results(self):
        p,d,o = fixture()
        state = create_initial_state("Dynamic Planning", p)
        state["planning_input"] = {"demands": [], "orders": o, "start_date": START, "horizon_days": 14}
        r = build_supply_chain_graph().invoke(state)
        self.assertTrue(r["workflow_error"])
        self.assertEqual([x["execution_status"] for x in r["execution_trace"]], ["FAILED", "SKIPPED", "SKIPPED", "SKIPPED"])

    def test_ai_context_is_dynamic_and_changes_with_orders(self):
        r = analyze_planning_bundle(build_planning_demo())["details"]["SEN-001"]
        c = build_assistant_context("single", single_result=r)
        self.assertEqual(c["payload"]["projection"]["total_unmet_demand"], 80)
        r2 = analyze_planning_bundle(build_planning_demo(scenario="正常到货"))["details"]["SEN-001"]
        c2 = build_assistant_context("single", single_result=r2)
        self.assertNotEqual(c["context_id"], c2["context_id"])
        class FakeLlm:
            def invoke(self, prompt):
                return '{"explanation_keys":["shortage","timing"]}'
        with patch("assistant._create_ollama_llm", return_value=FakeLlm()):
            answer = answer_supply_chain_question("为什么缺料？", c)
        self.assertIn("80", answer)
        with patch("assistant._create_ollama_llm") as llm:
            llm.return_value.invoke.return_value = "已下单100件"
            with self.assertRaises(AssistantUnavailableError):
                answer_supply_chain_question("怎么办？", c)

    def test_dynamic_page_interactions_and_legacy_round_trip(self):
        app = AppTest.from_file(ROOT / "app.py", default_timeout=30).run()
        app.get_by_key("analysis_mode").set_value("动态物料计划").run()
        app.get_by_key("planning_data_pack").select("公式核验 · 5种物料").run()
        app.get_by_key("planning_horizon").set_value(14).run()
        app.get_by_key("planning_scenario").select("供应商交付延迟").run()
        app.get_by_key("load_planning_demo").click().run()
        app.get_by_key("run_planning_analysis").click().run()
        self.assertEqual([e.value for e in app.exception], [])
        self.assertEqual(app.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 80)
        app.get_by_key("planning_scenario").select("正常到货").run()
        app.get_by_key("load_planning_demo").click().run()
        app.get_by_key("run_planning_analysis").click().run()
        self.assertEqual(app.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 0)
        app.get_by_key("planning_sku").select("BRK-004").run()
        self.assertEqual(app.session_state["planning_selected_result"]["selected_part"]["part_code"], "BRK-004")
        self.assertIsNotNone(app.get_by_key("download_planning_pdf"))
        self.assertIsNotNone(app.get_by_key("download_planning_template"))
        app.get_by_key("analysis_mode").set_value("单零部件情景分析").run()
        self.assertEqual(app.session_state["last_result"]["inventory_metrics"]["stockout_gap"], 0)
        self.assertEqual([e.value for e in app.exception], [])


if __name__ == "__main__":
    unittest.main()
