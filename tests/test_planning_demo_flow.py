import copy
import io
import sys
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "agent"))
from planning_data import (INTERVIEW_SCENARIOS, PlanningInputError, build_interview_demo,
                           validate_planning_bundle, analyze_planning_bundle,
                           planning_xlsx_bytes, read_planning_xlsx, normal_arrival_baseline)
from planning_workspace import analysis_totals, fingerprint, new_workspace
from workflow_runtime import WorkflowContractError
import planning_ui


class InterviewDataTests(unittest.TestCase):
    def projection(self, scenario="正常到货", horizon=14):
        return analyze_planning_bundle(build_interview_demo(horizon_days=horizon, scenario=scenario))["details"]["SEN-001"]["inventory_projection"]

    def test_four_scenarios_five_skus_same_stock_and_traceable_changes(self):
        base = build_interview_demo(horizon_days=14)
        for scenario in INTERVIEW_SCENARIOS:
            b = build_interview_demo(horizon_days=14, scenario=scenario)
            self.assertEqual(len(b["parts"]), 5)
            self.assertEqual(len(b["demands"]), 70)
            self.assertEqual(len(b["orders"]), 6)
            self.assertEqual([(p["part_code"], p["current_stock"]) for p in b["parts"]], [(p["part_code"], p["current_stock"]) for p in base["parts"]])
            self.assertEqual([o["open_quantity"] for o in b["orders"]], [o["open_quantity"] for o in base["orders"]])
            if scenario != "生产需求激增":
                self.assertEqual(b["demands"], base["demands"])
        delayed = build_interview_demo(horizon_days=14, scenario="供应商交付延迟")
        self.assertEqual(delayed["orders"][0]["arrival_date"], "2026-10-23")
        self.assertEqual(delayed["orders"][0]["original_arrival_date"], "2026-10-19")
        self.assertEqual(delayed["orders"][0]["status"], "CONFIRMED")
        self.assertEqual(delayed["orders"][0]["delivery_status"], "DELAYED")

    def test_14_day_business_acceptance_all_scenarios(self):
        expected = {"正常到货": (0, 180, None), "供应商交付延迟": (80, 260, "2026-10-22"),
                    "生产需求激增": (140, 0, "2026-10-24"), "采购订单待确认": (320, 0, "2026-10-22")}
        for scenario, values in expected.items():
            with self.subTest(scenario=scenario):
                r = self.projection(scenario)
                self.assertEqual((r["total_unmet_demand"], r["final_available"], r["first_shortage_date"]), values)
                self.assertEqual(len(r["daily"]), 14)

    def test_30_days_are_fully_simulated_and_not_inferred_from_14(self):
        for scenario, expected in zip(INTERVIEW_SCENARIOS, [1100, 1100, 1420, 1600]):
            r = self.projection(scenario, 30)
            self.assertEqual(len(r["daily"]), 30)
            self.assertEqual(r["daily"][-1]["date"], "2026-11-10")
            self.assertEqual(r["total_unmet_demand"], expected)
        self.assertEqual(self.projection(horizon=30)["first_shortage_date"], "2026-10-28")
        self.assertIsNone(self.projection(horizon=14)["first_shortage_date"])

    def test_order_date_edit_changes_real_inventory(self):
        b = build_interview_demo(horizon_days=14)
        before = analyze_planning_bundle(b)["details"]["SEN-001"]["inventory_projection"]
        b["orders"][0]["arrival_date"] = "2026-10-23"
        after = analyze_planning_bundle(b)["details"]["SEN-001"]["inventory_projection"]
        self.assertEqual((before["total_unmet_demand"], after["total_unmet_demand"]), (0, 80))
        self.assertEqual((before["daily"][7]["confirmed_receipts"], after["daily"][7]["confirmed_receipts"]), (500, 0))
        self.assertEqual(after["daily"][11]["confirmed_receipts"], 500)
        self.assertEqual(after["confirmed_orders"][0]["delivery_status"], "DELAYED")

    def test_normal_baseline_recomputed_without_mutating_actual_pending_orders(self):
        b = build_interview_demo(horizon_days=14, scenario="采购订单待确认")
        reference = normal_arrival_baseline(b)
        self.assertEqual(b["orders"][0]["status"], "PENDING")
        self.assertEqual(b["demands"], reference["demands"])
        self.assertEqual(analyze_planning_bundle(reference)["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 0)
        self.assertEqual(analyze_planning_bundle(b)["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 320)

    def test_named_example_round_trip_and_excel_date_modification(self):
        b = build_interview_demo(horizon_days=14)
        workbook = load_workbook(io.BytesIO(planning_xlsx_bytes(b, named_sheets=True)))
        self.assertEqual(workbook.sheetnames, ["SKU_Master", "Demand_Plan", "Purchase_Orders", "Suppliers", "Settings"])
        workbook["Purchase_Orders"]["E2"] = date(2026, 10, 23)
        payload = io.BytesIO()
        workbook.save(payload)
        read = read_planning_xlsx(payload.getvalue())
        self.assertEqual(analyze_planning_bundle(read)["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 80)
        # All four examples also round-trip without changing calculated results.
        for scenario in INTERVIEW_SCENARIOS:
            b = build_interview_demo(horizon_days=14, scenario=scenario)
            r = read_planning_xlsx(planning_xlsx_bytes(b, named_sheets=True))
            self.assertEqual(analyze_planning_bundle(b)["summary"].to_dict("records"), analyze_planning_bundle(r)["summary"].to_dict("records"))

    def test_blank_template_has_headers_but_is_not_a_valid_example(self):
        payload = planning_xlsx_bytes(build_interview_demo(), named_sheets=True, blank=True)
        workbook = load_workbook(io.BytesIO(payload))
        self.assertEqual(workbook["SKU_Master"].max_row, 1)
        self.assertEqual(workbook["Demand_Plan"].max_row, 1)
        with self.assertRaises(PlanningInputError):
            read_planning_xlsx(payload)

    def test_invalid_inputs_are_explicit_input_errors(self):
        mutations = [lambda b:b["orders"][0].pop("status"), lambda b:b["orders"][0].update(status="INVALID"),
                     lambda b:b["orders"][0].update(arrival_date="2026-02-30"),
                     lambda b:b["orders"].append(copy.deepcopy(b["orders"][0])),
                     lambda b:b["demands"][0].update(quantity=-1), lambda b:b["demands"][0].update(part_code="UNKNOWN"),
                     lambda b:b["orders"][0].update(delivery_status="INVALID")]
        for mutation in mutations:
            b = build_interview_demo(horizon_days=14)
            mutation(b)
            with self.subTest(mutation=mutation), self.assertRaises(PlanningInputError):
                analyze_planning_bundle(b)

    def test_totals_are_derived_from_projections_and_workspace_copy_is_independent(self):
        b = build_interview_demo(horizon_days=14, scenario="供应商交付延迟")
        r = analyze_planning_bundle(b)
        self.assertEqual(analysis_totals(r), {"sku_count": 5, "shortage_sku_count": 1,
                                             "first_shortage_date": "2026-10-22", "total_unmet_demand": 80})
        w = new_workspace(b, "demo", ("delay", "2026-10-12", 14))
        self.assertEqual(fingerprint(b), fingerprint(w["bundle"]))
        w["bundle"]["demands"][0]["quantity"] += 1
        self.assertNotEqual(fingerprint(b), fingerprint(w["bundle"]))
        self.assertEqual(b["demands"][0]["quantity"], 80)


class DemoFlowTests(unittest.TestCase):
    def test_upload_callback_loads_actual_excel_and_preserves_previous_analysis(self):
        b = build_interview_demo(horizon_days=14)
        b["orders"][0]["arrival_date"] = "2026-10-23"
        payload = planning_xlsx_bytes(b, named_sheets=True)
        class Upload:
            name = "fictional_modified_example.xlsx"
            def getvalue(self):
                return payload
        old = {"last_successful_analysis": True}
        state = {"planning_upload": Upload(), "planning_scenario": "正常到货",
                 "planning_start": date(2026, 10, 12), "planning_horizon": 30,
                 "planning_analysis": old}
        with patch.object(planning_ui.st, "session_state", state):
            planning_ui._import_upload()
        self.assertEqual(state["planning_horizon"], 14)
        self.assertEqual(state["planning_analysis"], old)
        self.assertIn("上传文件", state["planning_workspace"]["source"])
        r = analyze_planning_bundle(state["planning_draft"])
        self.assertEqual(r["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 80)
        self.assertEqual(analysis_totals(r)["sku_count"], 5)

    def app(self, scenario="正常到货", horizon=14, analyze=True):
        a = AppTest.from_file(ROOT / "app.py", default_timeout=30).run()
        a.get_by_key("analysis_mode").set_value("动态物料计划").run()
        a.get_by_key("planning_data_pack").select("公式核验 · 5种物料").run()
        a.get_by_key("planning_horizon").set_value(horizon).run()
        a.get_by_key("planning_scenario").select(scenario).run()
        a.get_by_key("load_planning_demo").click().run()
        if analyze:
            a.get_by_key("run_planning_analysis").click().run()
        self.assertEqual([e.value for e in a.exception], [])
        return a

    def test_default_has_five_skus_and_explicit_run_button(self):
        a = self.app(horizon=30, analyze=False)
        self.assertEqual(len(a.session_state["planning_draft"]["parts"]), 5)
        self.assertFalse(a.get_by_key("run_planning_analysis").disabled)
        with self.assertRaises(KeyError):
            a.session_state["planning_analysis"]
        a.get_by_key("run_planning_analysis").click().run()
        self.assertEqual(analysis_totals(a.session_state["planning_analysis"])["total_unmet_demand"], 1100)
        self.assertEqual([e.value for e in a.exception], [])

    def test_scene_selection_requires_load_and_analysis_does_not_silently_change(self):
        a = self.app()
        before = a.session_state["planning_analysis"]
        a.get_by_key("planning_scenario").select("供应商交付延迟").run()
        self.assertTrue(a.get_by_key("run_planning_analysis").disabled)
        self.assertEqual(a.session_state["planning_analysis"], before)
        self.assertTrue(a.session_state["planning_results_dirty"])
        a.get_by_key("load_planning_demo").click().run()
        self.assertEqual(a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 0)
        a.get_by_key("run_planning_analysis").click().run()
        self.assertEqual(a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 80)
        self.assertFalse(a.session_state["planning_results_dirty"])

    def test_date_editor_persists_rerun_sku_and_mode_changes(self):
        a = self.app()
        revision = a.session_state["planning_workspace"]["revision"]
        key = f"planning_editor_orders_{revision}"
        a.session_state[key] = {"edited_rows": {0: {"arrival_date": "2026-10-23"}}, "added_rows": [], "deleted_rows": []}
        a.run()
        self.assertTrue(a.session_state["planning_results_dirty"])
        self.assertEqual(a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 0)
        a.get_by_key("planning_sku").select("BRG-002").run()
        self.assertEqual(a.session_state["planning_draft"]["orders"][0]["arrival_date"], "2026-10-23")
        a.get_by_key("analysis_mode").set_value("单零部件情景分析").run()
        a.get_by_key("analysis_mode").set_value("动态物料计划").run()
        a.get_by_key("planning_data_pack").select("公式核验 · 5种物料").run()
        self.assertEqual(a.session_state["planning_draft"]["orders"][0]["arrival_date"], "2026-10-23")
        a.get_by_key("run_planning_analysis").click().run()
        self.assertEqual(a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["total_unmet_demand"], 80)
        self.assertEqual(a.get_by_key("planning_sku").value, "BRG-002")
        self.assertEqual([e.value for e in a.exception], [])

    def test_demand_editor_updates_computation_only_after_run(self):
        a = self.app()
        revision = a.session_state["planning_workspace"]["revision"]
        a.session_state[f"planning_editor_demands_{revision}"] = {"edited_rows": {0: {"quantity": 120}}, "added_rows": [], "deleted_rows": []}
        a.run()
        self.assertTrue(a.session_state["planning_results_dirty"])
        a.get_by_key("run_planning_analysis").click().run()
        p = a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]
        self.assertEqual(p["daily"][0]["planned_demand"], 120)
        self.assertEqual(p["final_available"], 140)

    def test_failed_input_preserves_last_successful_result(self):
        a = self.app()
        revision = a.session_state["planning_workspace"]["revision"]
        a.session_state[f"planning_editor_orders_{revision}"] = {"edited_rows": {0: {"arrival_date": "invalid-date"}}, "added_rows": [], "deleted_rows": []}
        a.get_by_key("run_planning_analysis").click().run()
        self.assertTrue(any("输入校验失败" in e.value for e in a.error))
        self.assertEqual(a.session_state["planning_analysis"]["details"]["SEN-001"]["inventory_projection"]["final_available"], 180)
        self.assertTrue(a.session_state["planning_results_dirty"])
        self.assertEqual([e.value for e in a.exception], [])

    def test_internal_failure_has_traceback_and_is_not_labeled_input_error(self):
        a = self.app(analyze=False)
        with patch.object(planning_ui, "analyze_planning_bundle", side_effect=WorkflowContractError("inventory_projection missing")):
            a.get_by_key("run_planning_analysis").click().run()
        self.assertTrue(any("工作流异常" in e.value for e in a.error))
        self.assertFalse(any("输入校验失败" in e.value for e in a.error))
        self.assertIn("Traceback", a.session_state["planning_internal_trace"])
        self.assertEqual([e.value for e in a.exception], [])


if __name__ == "__main__":
    unittest.main()
